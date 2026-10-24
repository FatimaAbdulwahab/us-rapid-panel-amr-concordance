#!/usr/bin/env python3
"""
build_docx.py - Build the Word (.docx) version of the preprint.

PLAIN-LANGUAGE NOTE: Uses the same converted manuscript text as the editable doc
(tex2md.py output, which reads numbers.json, preprint.aux, tables/ and
references.tex), so the Word file says exactly what the PDF says. Figures are the
same PNGs. Equations become native Word equations. Output: preprint.docx
"""
import os, re, subprocess
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SEC = os.path.join(HERE, "docsec")
FIG = os.path.join(HERE, "figures")
FIGMAP = {"fig1_flow": None, "fig2_cohort": None, "fig3_forest": None, "fig4_carb_drug": None, "fig5_ctxm": None,
          "fig6_mic": None, "fig7_discordance": None, "fig8_robustness": None}
FONT = "Palatino Linotype"

# 1. regenerate sections with local figure paths
subprocess.run(["python3", os.path.join(HERE, "tex2md.py"),
                "{" + ",".join(f'"{k}":"{os.path.join(FIG, k + ".png")}"' for k in FIGMAP) + "}"],
               check=True, stdout=subprocess.DEVNULL)

parts = []
front = open(os.path.join(SEC, "00_front.md")).read()
m = re.match(r"\*\*(.*?)\*\*\n\n(.*?)\n\n(## Abstract.*)", front, re.S)
title, author, rest = m.group(1), m.group(2), m.group(3)
parts.append("---\ntitle: \"" + title.replace('"', '\\"') + "\"\n---\n")
parts.append(author.replace("Correspondence:", "\n\nCorrespondence:") + "\n\n" + rest)
for f in ["01.md", "02.md", "03.md", "04.md", "05.md", "06_refs.md", "07_appendix.md"]:
    parts.append(open(os.path.join(SEC, f)).read())
md = "\n\n".join(parts)
# latex code blocks -> display math
md = re.sub(r"```latex\n(.*?)\n```", lambda m: "$$" + m.group(1).strip() + "$$", md, flags=re.S)
# image lines: give pandoc a width; the bold caption paragraph follows each image
md = re.sub(r"!\[(Fig\. \d+)\]\(([^)]+)\)", r"![](\2){width=6.5in}", md)
src = os.path.join(HERE, "preprint_docx.md")
open(src, "w").write(md)

out = os.path.join(HERE, "preprint.docx")
subprocess.run(["pandoc", src, "-f", "markdown+pipe_tables+tex_math_dollars+superscript", "-o", out,
                "--resource-path", HERE], check=True)

# 2. styling pass
d = Document(out)
sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(1.0))
sec.top_margin = sec.bottom_margin = Inches(1.0)

st0 = d.styles
_names = {s.name.lower(): s for s in st0}


class _St:
    def __getitem__(self, k):
        return _names[k.lower()]

    def __iter__(self):
        return iter(st0)


st = _St()
for name in ["Normal", "Body Text", "First Paragraph", "Compact", "Title", "Heading 1", "Heading 2", "Heading 3",
             "Heading 4", "Image Caption", "Table"]:
    if name.lower() not in _names:
        continue
    s = st[name]
    s.font.name = FONT
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    s.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
for name, size in [("Normal", 11), ("Body Text", 11), ("First Paragraph", 11), ("Compact", 11), ("Title", 16),
                   ("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 11), ("Heading 4", 11)]:
    if name.lower() in _names:
        st[name].font.size = Pt(size)
for name in ["Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title"]:
    if name.lower() in _names:
        st[name].font.bold = True
        st[name].font.italic = name == "Heading 4"
        st[name].font.color.rgb = RGBColor(0x11, 0x11, 0x11)
for name in ["Body Text", "First Paragraph"]:
    if name.lower() in _names:
        pf = st[name].paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
st["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, val in kw.items():
        el = OxmlElement(f"w:{edge}")
        for k, v in val.items():
            el.set(qn(f"w:{k}"), str(v))
        borders.append(el)


TEXT_DXA = 9360  # 6.5 in of text width


def size_columns(t, text_dxa=TEXT_DXA):
    """Column widths proportional to content (longest line, capped), filling the text width."""
    ncols = len(t.columns)
    need = []
    for ci in range(ncols):
        longest = 4
        for row in t.rows:
            txt = row.cells[ci].text.strip()
            words = max((len(w) for w in txt.split()), default=0)
            mono = 1.35 if re.match(r"^(SAMN|bla|mec|van)", txt) else 1.0
            longest = max(longest, min(len(txt), 24), int((words + 2) * mono))
        need.append(longest)
    tot = sum(need)
    # Floor narrow columns at 640 dxa, then share the remaining width among the rest
    # in proportion to content, so the total always equals the text width.
    floor = 640
    fixed = set()
    while True:
        free = [i for i in range(ncols) if i not in fixed]
        room = text_dxa - floor * len(fixed)
        ftot = sum(need[i] for i in free) or 1
        widths = [floor if i in fixed else int(room * need[i] / ftot) for i in range(ncols)]
        newly = {i for i in free if widths[i] < floor}
        if not newly:
            break
        fixed |= newly
    widths[max(range(ncols), key=lambda i: widths[i])] += text_dxa - sum(widths)
    tbl = t._tbl
    tblPr = tbl.tblPr
    for el in tblPr.findall(qn("w:tblW")):
        tblPr.remove(el)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(text_dxa))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
             "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
             "tblDescription"]
    kids = list(tblPr)
    for k in kids:
        tblPr.remove(k)
    kids.sort(key=lambda e: ORDER.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in ORDER else 99)
    for k in kids:
        tblPr.append(k)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for row in t.rows:
        for ci, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            for el in tcPr.findall(qn("w:tcW")):
                tcPr.remove(el)
            cw = OxmlElement("w:tcW")
            cw.set(qn("w:w"), str(widths[ci]))
            cw.set(qn("w:type"), "dxa")
            tcPr.insert(0, cw)


def sect_pr(landscape):
    sp = OxmlElement("w:sectPr")
    pg = OxmlElement("w:pgSz")
    if landscape:
        pg.set(qn("w:w"), "15840"); pg.set(qn("w:h"), "12240"); pg.set(qn("w:orient"), "landscape")
    else:
        pg.set(qn("w:w"), "12240"); pg.set(qn("w:h"), "15840")
    mar = OxmlElement("w:pgMar")
    for k in ("top", "bottom", "left", "right"):
        mar.set(qn(f"w:{k}"), "1440")
    for k, v in (("header", "720"), ("footer", "720"), ("gutter", "0")):
        mar.set(qn(f"w:{k}"), v)
    sp.append(pg)
    sp.append(mar)
    return sp


def end_section_at(p_el, landscape):
    pPr = p_el.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_el.insert(0, pPr)
    pPr.append(sect_pr(landscape))


WIDE = 8
for t in d.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    wide = len(t.columns) >= WIDE
    size_columns(t, 12960 if wide else TEXT_DXA)
    if wide:
        cap = t._tbl.getprevious()  # "Table n." caption paragraph
        before = cap.getprevious()
        end_section_at(before, landscape=False)
        nxt = t._tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p") and "".join(nxt.itertext()).startswith("Note:"):
            end_section_at(nxt, landscape=True)
        else:
            newp = OxmlElement("w:p")
            t._tbl.addnext(newp)
            end_section_at(newp, landscape=True)
    ncols = len(t.columns)
    size = Pt(7) if ncols >= 10 else (Pt(8) if ncols >= 8 else Pt(9))
    for ri, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                for r in p.runs:
                    r.font.size = size
                    if ri == 0:
                        r.font.bold = True
            if ri == 0:
                set_cell_border(cell, top={"val": "double", "sz": 6, "color": "000000"},
                                bottom={"val": "single", "sz": 6, "color": "000000"})
            if ri == len(t.rows) - 1:
                set_cell_border(cell, bottom={"val": "double", "sz": 6, "color": "000000"})

# captions: the bold "Fig. n." / "Table n." paragraphs -> smaller, keep tables with their captions
for p in d.paragraphs:
    txt = p.text
    if re.match(r"^(Fig\.|Table) [IVXA\d]+\.", txt) or txt.startswith("Note:"):
        for r in p.runs:
            r.font.size = Pt(9.5)
        p.paragraph_format.space_after = Pt(8)
        if txt.startswith("Table"):
            p.paragraph_format.keep_with_next = True
    if txt.startswith("Correspondence:") or txt.startswith("Fatima Abdulwahab (ORCID"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
# images: centered, keep with caption
for p in d.paragraphs:
    if p._p.xpath(".//pic:pic"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True

# page numbers in footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
    if tag:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), tag)
        run._r.append(fc)
    else:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = text
        run._r.append(it)

d.core_properties.title = title
d.core_properties.author = "Fatima Abdulwahab"
d.save(out)
os.remove(src)
print("wrote", out)


# 3. schema clean-up pass: element order inside pPr / style / sectPr, numbering ids
import zipfile, shutil
PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers",
       "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct",
       "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
       "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl",
       "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
STY = ["name", "aliases", "basedOn", "next", "link", "autoRedefine", "hidden", "uiPriority", "semiHidden",
       "unhideWhenUsed", "qFormat", "locked", "personal", "personalCompose", "personalReply", "rsid", "pPr", "rPr",
       "tblPr", "trPr", "tcPr", "tblStylePr"]
RPR = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow",
       "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern", "position",
       "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
       "eastAsianLayout", "specVanish", "oMath"]
from lxml import etree
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
        "tcFitText", "vAlign", "hideMark"]
SETTINGS = ["writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime",
            "doNotDisplayPageBoundaries", "displayBackgroundShape", "printPostScriptOverText",
            "printFractionalCharacterWidth", "printFormsData", "embedTrueTypeFonts", "embedSystemFonts",
            "saveSubsetFonts", "saveFormsData", "mirrorMargins", "alignBordersAndEdges", "bordersDoNotSurroundHeader",
            "bordersDoNotSurroundFooter", "gutterAtTop", "hideSpellingErrors", "hideGrammaticalErrors",
            "activeWritingStyle", "proofState", "formsDesign", "attachedTemplate", "linkStyles",
            "stylePaneFormatFilter", "stylePaneSortMethod", "documentType", "mailMerge", "revisionView",
            "trackRevisions", "doNotTrackMoves", "doNotTrackFormatting", "documentProtection", "autoFormatOverride",
            "styleLockTheme", "styleLockQFSet", "defaultTabStop", "autoHyphenation", "consecutiveHyphenLimit",
            "hyphenationZone", "doNotHyphenateCaps", "showEnvelope", "summaryLength", "clickAndTypeStyle",
            "defaultTableStyle", "evenAndOddHeaders", "bookFoldRevPrinting", "bookFoldPrinting",
            "bookFoldPrintingSheets", "drawingGridHorizontalSpacing", "drawingGridVerticalSpacing",
            "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery",
            "doNotUseMarginsForDrawingGridOrigin", "drawingGridHorizontalOrigin", "drawingGridVerticalOrigin",
            "doNotShadeFormData", "noPunctuationKerning", "characterSpacingControl", "printTwoOnOne",
            "strictFirstAndLastChars", "noLineBreaksAfter", "noLineBreaksBefore", "savePreviewPicture",
            "doNotValidateAgainstSchema", "saveInvalidXml", "ignoreMixedContent", "alwaysShowPlaceholderText",
            "doNotDemarcateInvalidXml", "saveXmlDataOnly", "useXSLTransform", "saveThroughXslt", "showXMLTags",
            "alwaysMergeEmptyNamespace", "updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr", "compat",
            "docVars", "rsids", "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
            "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade", "captions",
            "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
            "decimalSymbol", "listSeparator"]


def reorder(parent, order, ns=None):
    ns = ns or W
    kids = list(parent)
    known = [k for k in kids if k.tag.startswith(ns) and k.tag[len(ns):] in order]
    if not known:
        return
    for k in known:
        parent.remove(k)
    known.sort(key=lambda e: order.index(e.tag[len(ns):]))
    rest = [k for k in list(parent)]
    for k in rest:
        parent.remove(k)
    for k in known + rest:
        parent.append(k)


tmp = out + ".tmp"
with zipfile.ZipFile(out) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in ("word/document.xml", "word/styles.xml", "word/numbering.xml", "word/settings.xml"):
            root = etree.fromstring(data)
            for p in root.iter(W + "pPr"):
                reorder(p, PPR)
            for r in root.iter(W + "rPr"):
                if r.text and r.text.strip():
                    r.text = None
                reorder(r, RPR)
            for tc in root.iter(W + "tcPr"):
                reorder(tc, TCPR)
            for r in root.iter(W + "rPr"):
                for ch in r:
                    if ch.tail and ch.tail.strip():
                        ch.tail = None
            for dp in root.iter(M + "dPr"):
                reorder(dp, ["begChr", "sepChr", "endChr", "grow", "shp", "ctrlPr"], ns=M)
            for mc in root.iter(M + "mcPr"):
                reorder(mc, ["count", "mcJc"], ns=M)
            if root.tag == W + "settings":
                kids = list(root)
                for k in kids:
                    root.remove(k)
                loc = lambda e: etree.QName(e).localname
                kids.sort(key=lambda e: SETTINGS.index(loc(e)) if loc(e) in SETTINGS else len(SETTINGS))
                for k in kids:
                    root.append(k)
            for st_el in root.iter(W + "style"):
                reorder(st_el, STY)
            for pm in root.iter(W + "pgMar"):
                for k, v in (("header", "720"), ("footer", "720"), ("gutter", "0")):
                    if pm.get(W + k) is None:
                        pm.set(W + k, v)
            for ns in root.iter(W + "nsid"):
                v = ns.get(W + "val")
                if v and len(v) != 8:
                    ns.set(W + "val", v.rjust(8, "0")[-8:])
            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        zout.writestr(item, data)
shutil.move(tmp, out)
print("schema pass done")
